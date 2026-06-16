#!/usr/bin/env python3
"""
extract_neysa_skus.py
=====================
Convert the Velocis pricing spreadsheet into neysa_skus.json (the input the
matcher expects). Run this whenever the pricing sheet changes.

Usage:
    python extract_neysa_skus.py Velocis_Pricing_Sheet.xlsx neysa_skus.json

It reads the "GPU and CPU VM" sheet, keeps only GPU rows (L4, L40S, H100 SXM,
H100 NVL, H200 SXM — A100 is intentionally excluded as Neysa does not sell it),
and records:
  - on-demand price as hourly INR for the whole node
  - committed tiers (1/6/12/24/36-month) as MONTHLY INR for the whole node

The matcher converts these to USD per GPU per hour using its own FX rate.
All Neysa VMs run Ubuntu, so os is set to "Ubuntu".
"""
import sys, json, re

try:
    import openpyxl
except ImportError:
    sys.exit("Run: pip install openpyxl")

GPU_TYPES = {"L4", "L40S", "H100 SXM", "H100 NVL", "H200 SXM"}  # A100 excluded
TIERS = ["1mo", "6mo", "12mo", "24mo", "36mo"]


def num(x):
    if x is None:
        return None
    if isinstance(x, (int, float)):
        return float(x)
    s = re.sub(r"[^0-9.]", "", str(x))
    return float(s) if s else None


def find_gpu_sheet(wb):
    for name in wb.sheetnames:
        if "gpu" in name.lower() and "vm" in name.lower():
            return wb[name]
    return wb[wb.sheetnames[0]]


def main():
    src = sys.argv[1] if len(sys.argv) > 1 else "Velocis_Pricing_Sheet.xlsx"
    out = sys.argv[2] if len(sys.argv) > 2 else "neysa_skus.json"

    wb = openpyxl.load_workbook(src, data_only=True)
    ws = find_gpu_sheet(wb)

    skus = []
    for r in ws.iter_rows(values_only=True):
        if not r or len(r) < 12:
            continue
        cat, gpucat, sku = r[0], r[1], r[2]
        od = r[6]
        monthly_vals = r[7:12]   # 1,6,12,24,36-month columns
        if not gpucat or str(gpucat).strip() not in GPU_TYPES:
            continue
        m = re.match(r"\s*(\d+)\s*x", str(sku))
        count = int(m.group(1)) if m else 1
        monthly = {t: num(v) for t, v in zip(TIERS, monthly_vals) if num(v)}
        if not monthly:
            continue
        skus.append({
            "service": re.sub(r"\s+", " ", str(cat)).strip(),
            "gpu_type": str(gpucat).strip(),
            "gpu_count": count,
            "label": re.sub(r"\s+", " ", str(sku)).strip(),
            "on_demand_inr_hr": num(od),     # None for bare-metal (NA)
            "monthly_inr": monthly,
            "os": "Ubuntu",
        })

    with open(out, "w") as f:
        json.dump(skus, f, indent=2)
    print(f"Wrote {out} with {len(skus)} SKUs")


if __name__ == "__main__":
    main()
